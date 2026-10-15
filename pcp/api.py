"""HTTP API over the analytics engine, plus the static web app.

Prototype storage: uploaded datasets are kept in memory for the life of the server process.
"""
from __future__ import annotations

import io
import json
import pathlib
import uuid
from typing import Any

import numpy as np
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from . import report as report_builder
from .engine import attribution, currency, data, optimize, performance, risk, spending, stats
from .engine import mandate as md

ROOT = pathlib.Path(__file__).resolve().parent.parent
SAMPLES = {
    "etfs": ("course_etfs.csv", "Six ETFs, daily prices 2012-2017"),
    "indices": ("index_returns_monthly.csv", "Six asset-class indices, monthly returns 2000-2017"),
    "fx": ("uk_equity_fund_gbpusd.csv", "UK equity fund in GBP with the GBP/USD rate, daily (for the currency step)"),
}
MAX_UPLOAD = 25 * 1024 * 1024

app = FastAPI(title="Portfolio Construction Platform", version="0.1.0")
STORE: dict[str, dict] = {}


# ------------------------------------------------------------------ helpers
def clean(obj: Any) -> Any:
    """Make engine output JSON-safe (numpy types, NaN -> null)."""
    if isinstance(obj, dict):
        return {k: clean(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [clean(v) for v in obj]
    if isinstance(obj, np.ndarray):
        return clean(obj.tolist())
    if isinstance(obj, (np.floating, float)):
        return float(obj) if np.isfinite(obj) else None
    if isinstance(obj, np.integer):
        return int(obj)
    return obj


@app.exception_handler(ValueError)
async def value_error(_, exc: ValueError):
    return JSONResponse(status_code=422, content={"detail": str(exc)})


def get(dataset_id: str) -> data.Dataset:
    if dataset_id not in STORE:
        raise HTTPException(404, "Dataset not found. It may have been cleared when the server restarted; upload it again.")
    return STORE[dataset_id]["dataset"]


def summary(dataset_id: str) -> dict:
    rec = STORE[dataset_id]
    ds: data.Dataset = rec["dataset"]
    return {"id": dataset_id, "filename": rec["filename"], "assets": ds.assets, "names": ds.names,
            "kind": ds.kind, "frequency": ds.frequency, "periods_per_year": ds.periods_per_year,
            "observations": int(ds.returns.shape[0]), "start": str(ds.dates[0].date()),
            "end": str(ds.dates[-1].date()), "warnings": ds.warnings}


def register(content: bytes, filename: str, kind: str = "auto", frequency: str = "auto") -> dict:
    ds = data.load_table(content, filename, kind=kind, frequency=frequency)
    dataset_id = uuid.uuid4().hex[:12]
    STORE[dataset_id] = {"dataset": ds, "filename": filename}
    return summary(dataset_id)


def window(ds: data.Dataset, start: str | None, end: str | None):
    mask = data.slice_dates(ds, start, end)
    if mask.sum() < 12:
        raise ValueError("The chosen date range has fewer than 12 observations.")
    return ds.returns[mask], ds.dates[mask]


def resolve(ds: data.Dataset, body) -> tuple[np.ndarray, Any, dict]:
    """Returns in the analysis window, their dates, and the resolved assumptions."""
    r, dates = window(ds, body.start, body.end)
    return r, dates, stats.resolve_assumptions(r, ds.periods_per_year, body.assumptions, rf=body.rf)


def thin(n: int, limit: int = 500) -> np.ndarray:
    """Indices that downsample a long series for charting, always keeping the last point."""
    if n <= limit:
        return np.arange(n)
    idx = np.unique(np.linspace(0, n - 1, limit).round().astype(int))
    return idx


def check_weights(weights: list[float], n: int) -> np.ndarray:
    w = np.asarray(weights, dtype=float)
    if w.shape != (n,):
        raise ValueError(f"Expected {n} weights, one per asset.")
    if abs(w.sum()) < 1e-9:
        raise ValueError("Weights sum to zero.")
    return w / w.sum()


def mandate_of(ds: data.Dataset, body) -> md.Mandate:
    return md.parse(getattr(body, "mandate", None), ds.n_assets)


def weights_for(weights, n: int, m: md.Mandate) -> np.ndarray:
    """Fully invested mandates rescale weights to 100%; others keep them, with the rest in cash."""
    w = np.asarray(weights, dtype=float)
    if w.shape != (n,):
        raise ValueError(f"Expected {n} weights, one per asset.")
    if not m.full_investment:
        return w
    if abs(w.sum()) < 1e-9:
        raise ValueError("Weights sum to zero.")
    return w / w.sum()


def cash_line(m: md.Mandate, w: np.ndarray, rf: float) -> float:
    """Annual return from the cash line net of borrowing spread and short costs."""
    return m.cash_return(w, rf) - m.short_drag(w)


AXES = {None: ("Risk: annualised volatility (σ)", "Expected annual return"),
        "benchmark": ("Tracking error", "Active return vs benchmark"),
        "surplus": ("Surplus risk", "Surplus return")}
KEY_NAMES = {None: ("Minimum variance", "Maximum Sharpe"), "benchmark": ("Minimum tracking error", "Maximum information ratio"),
             "surplus": ("Minimum surplus risk", "Best surplus return per unit of risk")}


# ------------------------------------------------------------------ request models
class Base(BaseModel):
    assumptions: dict | None = None
    mandate: dict | None = None
    rf: float = 0.0
    start: str | None = None
    end: str | None = None


class OptimiseIn(Base):
    objective: str = "max_sharpe"
    target: float | None = None
    constraints: dict | None = None
    beta: float = Field(0.95, ge=0.5, lt=1)


class FrontierIn(Base):
    constraints: dict | None = None
    points: int = Field(40, ge=5, le=200)


class Portfolio(BaseModel):
    name: str
    weights: list[float]


class BacktestIn(Base):
    portfolios: list[Portfolio]
    benchmark: str | None = None
    rebalance: str = "every_period"
    cost_bps: float = Field(0.0, ge=0, le=500)


class StressIn(Base):
    weights: list[float]
    windows: list[dict] | None = None
    shock_window: dict | None = None


class SimulateIn(Base):
    weights: list[float]
    years: float = Field(5.0, gt=0, le=50)
    trials: int = Field(5000, ge=100, le=100000)
    method: str = "normal"
    seed: int = 42
    df: int = Field(5, ge=3, le=60)


class ReportIn(BacktestIn):
    constraints: dict | None = None


class AttributionIn(Base):
    weights: list[float]
    benchmark_weights: list[float] | None = None


class CurrencyIn(Base):
    asset: int
    currency: int
    hedge_ratio: float = Field(0.5, ge=0, le=1.5)
    base_rate: float = 0.0
    foreign_rate: float = 0.0


class SpendingIn(Base):
    weights: list[float]
    years: int = Field(30, ge=1, le=100)
    trials: int = Field(5000, ge=100, le=50000)
    seed: int = 42
    rule: str = "smoothed"
    rate: float = Field(0.04, ge=0, lt=1)
    smoothing_years: int = Field(3, ge=1, le=20)
    hybrid_weight: float = Field(0.7, ge=0, le=1)
    inflation: float = 0.025
    inflation_vol: float = Field(0.01, ge=0)
    contribution: float = Field(0.0, ge=0)
    fee: float = Field(0.0, ge=0, lt=0.2)
    compare_rates: list[float] | None = None


class ReportDocIn(ReportIn):
    title: str = "Strategic asset allocation review"
    client: str | None = None
    focus: str | None = None
    simulation: dict | None = None
    spending: dict | None = None
    stress_windows: list[dict] | None = None


# ------------------------------------------------------------------ datasets
@app.get("/api/health")
def health():
    return {"status": "ok", "datasets": len(STORE), "version": app.version}


DC_FILE, DC_META = ROOT / "data" / "dc_plan" / "dc_plan_daily.csv", ROOT / "data" / "dc_plan" / "dc_plan_meta.json"


@app.get("/api/samples")
def samples():
    out = []
    if DC_FILE.exists():                      # the weekly-updated DC plan file, when it has been built
        meta = json.loads(DC_META.read_text(encoding="utf-8")) if DC_META.exists() else {}
        out.append({"key": "dc_plan", "description": f"My DC plan: {len(meta.get('assets', []))} asset classes, daily to {meta.get('end', '?')}"})
    return out + [{"key": k, "description": d} for k, (_, d) in SAMPLES.items()]


@app.get("/api/dc-plan")
def dc_plan_status():
    """When the DC plan data was last refreshed, and any warnings from that update."""
    if not DC_META.exists():
        raise HTTPException(404, "The DC plan data has not been downloaded yet. Run tools/update_dc_data.py.")
    return json.loads(DC_META.read_text(encoding="utf-8"))


@app.post("/api/datasets/sample/{key}")
def load_sample(key: str):
    if key == "dc_plan" and DC_FILE.exists():
        return register(DC_FILE.read_bytes(), DC_FILE.name)
    if key not in SAMPLES:
        raise HTTPException(404, "Unknown sample.")
    name = SAMPLES[key][0]
    return register((ROOT / "sample_data" / name).read_bytes(), name)


@app.post("/api/datasets")
async def upload(file: UploadFile = File(...), kind: str = Form("auto"), frequency: str = Form("auto")):
    content = await file.read()
    if len(content) > MAX_UPLOAD:
        raise HTTPException(413, "File is larger than 25 MB.")
    if not (file.filename or "").lower().endswith((".csv", ".txt", ".xlsx", ".xlsm", ".xls")):
        raise ValueError("Upload a CSV or Excel file.")
    return register(content, file.filename, kind, frequency)


@app.get("/api/datasets/{dataset_id}")
def dataset(dataset_id: str):
    get(dataset_id)
    return summary(dataset_id)


# ------------------------------------------------------------------ analyses
@app.post("/api/datasets/{dataset_id}/statistics")
def statistics(dataset_id: str, body: Base):
    ds = get(dataset_id)
    r, dates = window(ds, body.start, body.end)
    spec = body.assumptions or {}
    a = stats.resolve_assumptions(r, ds.periods_per_year, spec, rf=body.rf)
    hist = stats.resolve_assumptions(r, ds.periods_per_year, {"method": spec.get("method", "arithmetic"),
                                                              "covariance": spec.get("covariance")})
    idx = thin(len(r) + 1)
    growth = stats.growth(r)[idx]
    labels = [str(dates[0].date())] + [str(d.date()) for d in dates]
    return clean({
        "assets": stats.asset_statistics(r, ds.periods_per_year, body.rf, spec.get("method", "arithmetic")),
        "assumptions": {k: a.get(k) for k in ("mu", "vol", "corr", "cov", "shrinkage", "notes", "prior_mu")},
        "historical": {k: hist[k] for k in ("mu", "vol", "corr")},
        "growth": {"dates": [labels[i] for i in idx], "values": growth.T},
        "histograms": [stats.histogram(r[:, i]) for i in range(ds.n_assets)],
        "observations": int(len(r)), "start": str(dates[0].date()), "end": str(dates[-1].date()),
    })


@app.post("/api/datasets/{dataset_id}/optimise")
def optimise(dataset_id: str, body: OptimiseIn):
    ds = get(dataset_id)
    r, _, a = resolve(ds, body)
    return clean(optimize.optimise(a["mu"], a["cov"], body.objective, rf=body.rf, target=body.target,
                                   constraints=body.constraints, scenarios=r, beta=body.beta,
                                   mandate=mandate_of(ds, body), ppy=ds.periods_per_year))


@app.post("/api/datasets/{dataset_id}/frontier")
def frontier(dataset_id: str, body: FrontierIn):
    ds = get(dataset_id)
    r, _, a = resolve(ds, body)
    mu, cov, m, n = a["mu"], a["cov"], mandate_of(ds, body), ds.n_assets
    front = optimize.efficient_frontier(mu, cov, rf=body.rf, constraints=body.constraints, points=body.points, mandate=m)
    low, ratio = KEY_NAMES[m.relative]
    key = {low: front["min_variance"], "Maximum return": front["max_return"]}
    notes = []
    # under a long-only benchmark mandate the information ratio peaks at tiny active bets, so it is not offered
    candidates = ([] if m.relative == "benchmark" else [(ratio, "max_sharpe")]) + \
        ([("Risk parity", "risk_parity")] if m.kind == "long_only" and m.reference is None else [])
    for label, objective in candidates:
        try:
            key[label] = optimize.optimise(mu, cov, objective, rf=body.rf, constraints=body.constraints, mandate=m)
        except optimize.Infeasible as exc:
            notes.append(f"{label}: {exc}")
    if m.kind != "market_neutral":
        key["Equal weight"] = optimize.describe(np.full(n, 1 / n), mu, cov, body.rf, m, normalise=True)
    if m.risk_budget is not None:
        key["Maximum return"]["note"] = f"Uses the full {m.risk_budget:.1%} risk budget or less."
    if m.relative == "benchmark":
        key["Benchmark"] = optimize.describe(m.reference, mu, cov, body.rf, m, normalise=False)
    elif m.relative == "surplus":
        key["Liability hedge"] = optimize.describe(np.eye(n)[m.liability], mu, cov, body.rf, m, normalise=False)
    x, y = AXES[m.relative]
    return clean({"points": front["points"], "portfolios": key, "notes": notes, "axes": {"x": x, "y": y},
                  "mandate": {"kind": m.kind, "label": m.label, "relative": m.relative, "full_investment": m.full_investment},
                  "assets": [{**optimize.describe(np.eye(n)[i], mu, cov, body.rf, m, normalise=False), "exp_return": mu[i],
                              "volatility": a["vol"][i]} for i in range(n)]})


@app.post("/api/datasets/{dataset_id}/backtest")
def backtest(dataset_id: str, body: BacktestIn):
    ds = get(dataset_id)
    r, dates = window(ds, body.start, body.end)
    ppy, m = ds.periods_per_year, mandate_of(ds, body)
    series = {}
    for p in body.portfolios:
        series[p.name] = performance.portfolio_returns(
            r, weights_for(p.weights, ds.n_assets, m), ppy=ppy, rebalance=body.rebalance, cost_bps=body.cost_bps,
            normalise=m.full_investment, cash_rate=body.rf, borrow_spread=m.borrow_spread, short_cost=m.short_cost)
    if body.benchmark and body.benchmark not in series:
        raise ValueError("The benchmark must be one of the portfolios in the request.")
    bench = series[body.benchmark][0] if body.benchmark else None
    idx = thin(len(r) + 1)
    labels = [str(dates[0].date())] + [str(d.date()) for d in dates]
    out = []
    for name, (s, turnover) in series.items():
        wealth, dd = performance.wealth_and_drawdown(s)
        m = performance.metrics(s, ppy=ppy, rf=body.rf, benchmark=bench if name != body.benchmark else None)
        out.append({"name": name, "metrics": {**m, "turnover": turnover}, "wealth": wealth[idx], "drawdown": dd[idx]})
    return clean({"dates": [labels[i] for i in idx], "portfolios": out, "observations": int(len(r)),
                  "start": str(dates[0].date()), "end": str(dates[-1].date())})


@app.post("/api/datasets/{dataset_id}/stress")
def stress(dataset_id: str, body: StressIn):
    ds = get(dataset_id)
    m = mandate_of(ds, body)
    w = weights_for(body.weights, ds.n_assets, m)
    ppy, extra = ds.periods_per_year, cash_line(m, w, body.rf)
    windows = body.windows or risk.suggest_windows(ds.dates, ds.returns @ w + extra / ppy, ppy)
    result = {"windows": risk.window_analysis(ds.returns, ds.dates, w, windows, ppy=ppy, rf=body.rf, extra_return=extra)}
    if body.shock_window:
        r, _ = window(ds, body.start, body.end)
        base = stats.resolve_assumptions(r, ppy, body.assumptions, rf=body.rf)["mu"]
        shock = risk.window_analysis(ds.returns, ds.dates, w, [body.shock_window], ppy=ppy)[0]
        if "error" in shock:
            raise ValueError(shock["error"])
        result["sensitivity"] = {**risk.shock_sensitivity(w, base, shock["asset_return"]),
                                 "base_mu": base, "shock_mu": shock["asset_return"], "window": body.shock_window}
    return clean(result)


@app.post("/api/datasets/{dataset_id}/simulate")
def simulate(dataset_id: str, body: SimulateIn):
    ds = get(dataset_id)
    r, _, a = resolve(ds, body)
    m = mandate_of(ds, body)
    w = weights_for(body.weights, ds.n_assets, m)
    return clean(risk.monte_carlo(w, a["mu"], a["cov"], ppy=ds.periods_per_year, years=body.years, trials=body.trials,
                                  method=body.method, seed=body.seed, history=r, df=body.df, rf=body.rf,
                                  extra_return=cash_line(m, w, body.rf)))


@app.post("/api/datasets/{dataset_id}/report.xlsx")
def report(dataset_id: str, body: ReportIn):
    """Audit workbook: inputs, assumptions, portfolios and backtest results as plain values."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    ds = get(dataset_id)
    r, dates = window(ds, body.start, body.end)
    ppy = ds.periods_per_year
    a = stats.resolve_assumptions(r, ppy, body.assumptions, rf=body.rf)
    m = mandate_of(ds, body)
    wb = Workbook()
    bold = Font(name="Arial", bold=True)

    def sheet(title, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        for row in rows:
            ws.append([None if isinstance(v, float) and not np.isfinite(v) else v for v in row])
        for cell in ws[1]:
            cell.font = bold
        ws.column_dimensions["A"].width = 34
        return ws

    info = summary(dataset_id)
    sheet("Inputs", [["Item", "Value"], ["Source file", info["filename"]], ["Data type", ds.kind],
                     ["Frequency", ds.frequency], ["Analysis start", str(dates[0].date())],
                     ["Analysis end", str(dates[-1].date())], ["Observations", int(len(r))],
                     ["Risk-free rate (annual)", body.rf], ["Assumption source", (body.assumptions or {}).get("source", "historical")],
                     ["Return multiplier", (body.assumptions or {}).get("return_multiplier", 1.0)],
                     ["Volatility multiplier", (body.assumptions or {}).get("vol_multiplier", 1.0)],
                     ["Mandate", m.label], ["Net exposure range", f"{m.net_min:.0%} to {m.net_max:.0%}"],
                     ["Gross exposure limit", "none" if m.gross_max is None else f"{m.gross_max:.0%}"],
                     ["Rebalancing", body.rebalance], ["Trading cost (bps)", body.cost_bps],
                     ["Note", "Analytics output for review. Not investment advice."]], first=True)
    sheet("Assumptions", [["Asset", "Name", "Expected return", "Volatility", "Min weight", "Max weight"]] + [
        [ds.assets[i], ds.names[i], float(a["mu"][i]), float(a["vol"][i]),
         ((body.constraints or {}).get("lower") or [0.0] * ds.n_assets)[i],
         ((body.constraints or {}).get("upper") or [1.0] * ds.n_assets)[i]] for i in range(ds.n_assets)])
    sheet("Correlation", [[""] + ds.assets] + [[ds.assets[i]] + [float(x) for x in a["corr"][i]] for i in range(ds.n_assets)])
    sheet("Covariance", [[""] + ds.assets] + [[ds.assets[i]] + [float(x) for x in a["cov"][i]] for i in range(ds.n_assets)])
    rows = [["Portfolio"] + ds.assets + ["Cash", "Gross exposure", "Expected return", "Volatility", "Sharpe", "Tracking / surplus risk"]]
    for p in body.portfolios:
        d = optimize.describe(weights_for(p.weights, ds.n_assets, m), a["mu"], a["cov"], body.rf, m)
        rows.append([p.name] + d["weights"] + [d["cash"], d["gross"], d["exp_return"], d["volatility"], d["sharpe"], d.get("tracking_error")])
    sheet("Portfolios", rows)
    bt = backtest(dataset_id, body)
    keys = ["ann_return", "volatility", "sharpe", "sortino", "max_drawdown", "cumulative", "worst", "best",
            "var_95", "cvar_95", "turnover", "active_return", "tracking_error", "information_ratio", "beta"]
    sheet("Backtest", [["Portfolio"] + keys] + [[p["name"]] + [p["metrics"].get(k) for k in keys] for p in bt["portfolios"]])
    sheet("Backtest paths", [["Date"] + [p["name"] for p in bt["portfolios"]]] + [
        [d] + [p["wealth"][i] for p in bt["portfolios"]] for i, d in enumerate(bt["dates"])])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="portfolio_analysis.xlsx"'})

RULE_LABELS = {"percent_of_value": "percentage of fund value", "smoothed": "percentage of the smoothed average value",
               "hybrid": "hybrid (Yale / Stanford)", "fixed_real": "fixed amount in today's money"}
SOURCE_LABELS = {"historical": "historical estimates from the uploaded data", "custom": "custom forecasts",
                 "blend": "a weighted blend of forecast sets", "black_litterman": "Black-Litterman equilibrium returns adjusted for views"}


@app.post("/api/datasets/{dataset_id}/attribution")
def attribute(dataset_id: str, body: AttributionIn):
    ds = get(dataset_id)
    r, dates, a = resolve(ds, body)
    m = mandate_of(ds, body)
    w = weights_for(body.weights, ds.n_assets, m)
    b = weights_for(body.benchmark_weights, ds.n_assets, m) if body.benchmark_weights else None
    labels, risk_share = list(ds.assets), optimize.describe(w, a["mu"], a["cov"], body.rf, m, normalise=False)["risk_contribution"]
    if abs(w.sum() - 1) > 1e-9 or (b is not None and abs(b.sum() - 1) > 1e-9):     # cash becomes its own line
        rf_p = (1 + body.rf) ** (1 / ds.periods_per_year) - 1
        r = np.column_stack([r, np.full(len(r), rf_p)])
        w = np.append(w, 1 - w.sum())
        b = np.append(b, 1 - b.sum()) if b is not None else None
        labels, risk_share = labels + ["Cash"], risk_share + [0.0]
    out = attribution.attribute(r, w, b)
    out["risk_contribution"] = risk_share
    return clean({**out, "labels": labels, "start": str(dates[0].date()), "end": str(dates[-1].date())})


@app.post("/api/datasets/{dataset_id}/currency")
def currency_hedge(dataset_id: str, body: CurrencyIn):
    ds = get(dataset_id)
    r, dates = window(ds, body.start, body.end)
    n = ds.n_assets
    if not (0 <= body.asset < n and 0 <= body.currency < n) or body.asset == body.currency:
        raise ValueError("Choose one column for the asset and a different one for the currency.")
    out = currency.hedge_analysis(r[:, body.asset], r[:, body.currency], ppy=ds.periods_per_year,
                                  hedge_ratio=body.hedge_ratio, base_rate=body.base_rate, foreign_rate=body.foreign_rate)
    idx = thin(len(r) + 1)
    labels = [str(dates[0].date())] + [str(d.date()) for d in dates]
    out["wealth"] = {k: np.asarray(v)[idx] for k, v in out["wealth"].items()}
    return clean({**out, "dates": [labels[i] for i in idx], "asset": ds.assets[body.asset], "currency": ds.assets[body.currency]})


@app.post("/api/datasets/{dataset_id}/spending")
def spending_policy(dataset_id: str, body: SpendingIn):
    ds = get(dataset_id)
    _, _, a = resolve(ds, body)
    m = mandate_of(ds, body)
    w = weights_for(body.weights, ds.n_assets, m)
    kw = dict(years=body.years, trials=body.trials, seed=body.seed, rule=body.rule, smoothing_years=body.smoothing_years,
              hybrid_weight=body.hybrid_weight, inflation=body.inflation, inflation_vol=body.inflation_vol,
              contribution=body.contribution, fee=body.fee, extra_return=cash_line(m, w, body.rf), normalise=m.full_investment)
    out = spending.simulate(w, a["mu"], a["cov"], rate=body.rate, **kw)
    out.update({"rule_label": RULE_LABELS[body.rule], "inflation": body.inflation, "contribution": body.contribution, "fee": body.fee})
    if body.compare_rates:
        out["comparison"] = spending.compare_rates(w, a["mu"], a["cov"], body.compare_rates[:12], **{**kw, "trials": min(body.trials, 3000)})
    return clean(out)


@app.post("/api/datasets/{dataset_id}/report.docx")
def report_docx(dataset_id: str, body: ReportDocIn):
    """Investment-committee report (Word) for the focus portfolio, with the other portfolios for comparison."""
    ds = get(dataset_id)
    if not body.portfolios:
        raise ValueError("Add at least one portfolio before creating the report.")
    r, dates, a = resolve(ds, body)
    n, ppy = ds.n_assets, ds.periods_per_year
    m = mandate_of(ds, body)
    names = [p.name for p in body.portfolios]
    focus = body.focus if body.focus in names else names[0]
    fw = weights_for(body.portfolios[names.index(focus)].weights, n, m)
    extra = cash_line(m, fw, body.rf)
    shared = {"assumptions", "mandate", "rf", "start", "end"}
    spec = body.assumptions or {}
    cons = body.constraints or {}
    sim = body.simulation or {}
    bench = body.benchmark if body.benchmark in names else None
    ctx = {
        "title": body.title, "client": body.client, "focus": focus, "rf": body.rf, "benchmark": bench,
        "mandate": {"label": m.label, "relative": m.relative, "full_investment": m.full_investment, "net_min": m.net_min,
                    "net_max": m.net_max, "gross_max": m.gross_max, "short_cost": m.short_cost, "borrow_spread": m.borrow_spread,
                    "active_limit": m.active_limit, "turnover_max": m.turnover_max, "funding_ratio": m.funding_ratio,
                    "liability": ds.assets[m.liability] if m.liability is not None else None},
        "axes": dict(zip(("x", "y"), AXES[m.relative])),
        "assets": ds.assets, "dataset": summary(dataset_id), "window": (str(dates[0].date()), str(dates[-1].date())),
        "mu": a["mu"].tolist(), "vol": a["vol"].tolist(), "notes": a["notes"],
        "assumption_label": SOURCE_LABELS.get(spec.get("source", "historical"), "historical estimates")
        + (f", returns x {spec['return_multiplier']:g} and risk x {spec.get('vol_multiplier', 1):g}"
           if spec.get("return_multiplier", 1) != 1 or spec.get("vol_multiplier", 1) != 1 else ""),
        "constraints": {"lower": cons.get("lower") or [(-1.0 if m.shorts else 0.0)] * n, "upper": cons.get("upper") or [max(1.0, m.net_max)] * n,
                        "groups": cons.get("groups") or []},
        "portfolios": [{"name": p.name, **{k: v for k, v in optimize.describe(weights_for(p.weights, n, m), a["mu"], a["cov"], body.rf, m).items()
                                              if k != "risk_contribution"}} for p in body.portfolios],
        "frontier": frontier(dataset_id, FrontierIn(**body.model_dump(include=shared | {"constraints"}))),
        "backtest": backtest(dataset_id, BacktestIn(**{**body.model_dump(include=shared | {"portfolios", "rebalance", "cost_bps"}), "benchmark": bench})),
        "rebalance_label": {"every_period": "rebalanced every period", "monthly": "rebalanced monthly", "quarterly": "rebalanced quarterly",
                            "annual": "rebalanced annually", "never": "never rebalanced (buy and hold)"}.get(body.rebalance, body.rebalance),
        "cost_bps": body.cost_bps,
        "stress": risk.window_analysis(ds.returns, ds.dates, fw, body.stress_windows or risk.suggest_windows(ds.dates, ds.returns @ fw + extra / ppy, ppy),
                                       ppy=ppy, rf=body.rf, extra_return=extra),
        "simulation": risk.monte_carlo(fw, a["mu"], a["cov"], ppy=ppy, years=float(sim.get("years", 5)), trials=int(sim.get("trials", 5000)),
                                       method=sim.get("method", "normal"), seed=int(sim.get("seed", 42)), history=r, rf=body.rf, extra_return=extra),
        "attribution": attribute(dataset_id, AttributionIn(**body.model_dump(include=shared), weights=fw.tolist(),
                                 benchmark_weights=body.portfolios[names.index(bench)].weights if bench and bench != focus else None)),
    }
    if body.spending:
        ctx["spending"] = spending_policy(dataset_id, SpendingIn(**{**body.model_dump(include=shared),
                                                                    **body.spending, "weights": fw.tolist()}))
    content = report_builder.build(ctx)
    return StreamingResponse(io.BytesIO(content), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                             headers={"Content-Disposition": 'attachment; filename="investment_committee_report.docx"'})


app.mount("/", StaticFiles(directory=ROOT / "web", html=True), name="web")
