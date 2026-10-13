"""Investment-committee report as a Word document, built from analysis results."""
from __future__ import annotations

import datetime
import io

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import matplotlib.ticker  # noqa: E402
import numpy as np  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

PALETTE = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7", "#e34948"]
GREY, INK = "#898781", "#0b0b0b"
pct = lambda v, d=1: "n/a" if v is None else f"{v * 100:.{d}f}%"
num = lambda v, d=2: "n/a" if v is None else f"{v:.{d}f}"
colour = lambda i: PALETTE[i] if i < len(PALETTE) else GREY


def _figure(width=6.4, height=3.3):
    fig, ax = plt.subplots(figsize=(width, height), dpi=170)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color("#c3c2b7")
    ax.grid(axis="y", color="#e1e0d9", linewidth=0.7)
    ax.set_axisbelow(True)
    ax.tick_params(colors=GREY, labelsize=8, length=0)
    return fig, ax


def _png(fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.tight_layout()
    fig.savefig(buf, format="png")
    plt.close(fig)
    buf.seek(0)
    return buf


def _pct_axis(ax, which="y", decimals=0):
    fmt = matplotlib.ticker.FuncFormatter(lambda v, _: f"{v * 100:.{decimals}f}%")
    (ax.yaxis if which == "y" else ax.xaxis).set_major_formatter(fmt)


def chart_frontier(ctx):
    fig, ax = _figure()
    pts = ctx["frontier"]["points"]
    ax.plot([p["volatility"] for p in pts], [p["exp_return"] for p in pts], color=PALETTE[0], linewidth=2.2, label="Efficient frontier")
    for i, a in enumerate(ctx["frontier"]["assets"]):
        ax.scatter(a["volatility"], a["exp_return"], marker="D", s=22, facecolor="white", edgecolor=INK, linewidth=1, zorder=3)
        ax.annotate(ctx["assets"][i], (a["volatility"], a["exp_return"]), textcoords="offset points", xytext=(5, -3), fontsize=7, color="#52514e")
    for k, p in enumerate(ctx["portfolios"]):
        ax.scatter(p["volatility"], p["exp_return"], s=46, color=colour(k), edgecolor="white", linewidth=1, zorder=4, label=p["name"])
    _pct_axis(ax, "y"), _pct_axis(ax, "x")
    ax.grid(axis="x", color="#e1e0d9", linewidth=0.7)
    ax.set_xlabel("Risk: annualised volatility", fontsize=8, color="#52514e")
    ax.set_ylabel("Expected annual return", fontsize=8, color="#52514e")
    ax.legend(fontsize=7, frameon=False, loc="best")
    return _png(fig)


def chart_weights(ctx):
    ports = ctx["portfolios"]
    fig, ax = _figure(height=0.55 + 0.42 * len(ports))
    left = np.zeros(len(ports))
    for i, asset in enumerate(ctx["assets"]):
        vals = np.array([p["weights"][i] for p in ports])
        ax.barh(range(len(ports)), vals, left=left, color=colour(i), edgecolor="white", linewidth=1, label=asset, height=0.62)
        for k, v in enumerate(vals):
            if v >= 0.07:
                ax.text(left[k] + v / 2, k, f"{v * 100:.0f}%", ha="center", va="center", fontsize=7,
                        color="white" if i in (0, 5, 6) else INK)
        left += vals
    ax.set_yticks(range(len(ports)), [p["name"] for p in ports], fontsize=8, color=INK)
    ax.invert_yaxis()
    ax.set_xlim(0, 1)
    ax.xaxis.set_visible(False)
    ax.grid(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_visible(False)
    ax.legend(fontsize=7, frameon=False, ncol=min(4, len(ctx["assets"])), loc="upper center", bbox_to_anchor=(0.5, -0.02))
    return _png(fig)


def chart_lines(dates, series, ylabel, percent=False, height=3.0):
    fig, ax = _figure(height=height)
    x = np.array(dates, dtype="datetime64[D]") if dates else None
    for k, (name, values) in enumerate(series):
        ax.plot(x if x is not None else range(len(values)), values, color=colour(k), linewidth=1.5, label=name)
    if percent:
        _pct_axis(ax, "y")
    ax.set_ylabel(ylabel, fontsize=8, color="#52514e")
    ax.legend(fontsize=7, frameon=False, loc="best")
    return _png(fig)


def chart_fan(values, xlabel, ylabel, per_year=1):
    fig, ax = _figure(height=3.0)
    v = np.array(values)                      # rows: 5, 25, 50, 75, 95
    x = np.arange(v.shape[1]) / per_year
    ax.fill_between(x, v[0], v[4], color=PALETTE[0], alpha=0.14, linewidth=0, label="5th to 95th percentile")
    ax.fill_between(x, v[1], v[3], color=PALETTE[0], alpha=0.28, linewidth=0, label="25th to 75th percentile")
    ax.plot(x, v[2], color=INK, linewidth=1.8, label="Median")
    ax.set_xlabel(xlabel, fontsize=8, color="#52514e")
    ax.set_ylabel(ylabel, fontsize=8, color="#52514e")
    ax.legend(fontsize=7, frameon=False, loc="upper left")
    return _png(fig)


# ------------------------------------------------------------------ document helpers
def _table(doc, head, rows, widths=None, bold_first=True):
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(head):
        cell = t.rows[0].cells[i]
        cell.text = str(h)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold, r.font.size = True, Pt(8.5)
            if i:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    r.font.bold = bool(bold_first and i == 0)
                if i:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


def _para(doc, text, bold_lead=None, size=10):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead + " ")
        r.bold, r.font.size = True, Pt(size)
    r = p.add_run(text)
    r.font.size = Pt(size)
    return p


def _bullets(doc, items):
    for text in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(text).font.size = Pt(10)


def build(ctx: dict) -> bytes:
    """ctx is assembled by the API from engine results; see pcp.api.report_docx."""
    doc = Document()
    for name in ("Normal", "Heading 1", "Heading 2", "Title", "List Bullet"):
        doc.styles[name].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    for name, size in (("Heading 1", 14), ("Heading 2", 11.5)):
        doc.styles[name].font.size, doc.styles[name].font.color.rgb = Pt(size), RGBColor(0x1F, 0x38, 0x64)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.8)

    focus = next(p for p in ctx["portfolios"] if p["name"] == ctx["focus"])
    bt = {p["name"]: p["metrics"] for p in ctx["backtest"]["portfolios"]}
    fm = bt[focus["name"]]
    assets = ctx["assets"]

    title = doc.add_heading(ctx["title"], level=0)
    title.runs[0].font.size = Pt(20)
    sub = doc.add_paragraph()
    sub.add_run(" · ".join(x for x in (ctx.get("client"), f"Prepared {datetime.date.today():%d %B %Y}",
                                         f"Data: {ctx['dataset']['filename']}") if x)).font.color.rgb = RGBColor(0x52, 0x51, 0x4E)

    # 1 summary
    doc.add_heading("1. Summary for the committee", level=1)
    top = sorted(zip(focus["weights"], assets), reverse=True)[:3]
    bullets = [
        f"Proposed allocation: {focus['name']}. Expected return {pct(focus['exp_return'])} a year at {pct(focus['volatility'])} "
        f"volatility (Sharpe ratio {num(focus['sharpe'])} at a {pct(ctx['rf'])} risk-free rate). "
        f"Largest holdings: {', '.join(f'{a} {pct(w, 0)}' for w, a in top if w > 0.005)}.",
        f"History: over {ctx['backtest']['start']} to {ctx['backtest']['end']} it would have returned {pct(fm['ann_return'])} a year "
        f"with {pct(fm['volatility'])} volatility and a worst peak-to-trough fall of {pct(fm['max_drawdown'])}.",
    ]
    if ctx.get("benchmark") and fm.get("tracking_error") is not None:
        bullets.append(f"Against {ctx['benchmark']}: {'ahead by' if fm['active_return'] >= 0 else 'behind by'} "
                       f"{pct(abs(fm['active_return']))} a year, tracking error {pct(fm['tracking_error'])}, "
                       f"information ratio {num(fm['information_ratio'])}.")
    ok = [w for w in ctx["stress"] if "error" not in w]
    if ok:
        worst = min(ok, key=lambda w: w["cumulative"])
        bullets.append(f"Stress: the hardest period in the data was {worst['name']} ({worst['start']} to {worst['end']}), "
                       f"with a cumulative return of {pct(worst['cumulative'])} and a drawdown of {pct(worst['max_drawdown'])}.")
    mc = ctx["simulation"]
    bullets.append(f"Outlook: over {mc['years']:.0f} years the simulated median return is {pct(mc['median_ann_return'])} a year; "
                   f"in 1 case in 20 it is {pct(mc['p05_ann_return'])} or lower. The chance of ending below the starting value is {pct(mc['prob_loss'])}.")
    sp = ctx.get("spending")
    if sp:
        bullets.append(f"Spending: with a {pct(sp['rate'])} payout under the {sp['rule_label']} rule, the fund keeps its "
                       f"inflation-adjusted value over {sp['years']} years in {pct(sp['prob_preserve'], 0)} of simulations.")
    _bullets(doc, bullets)

    # 2 data & assumptions
    doc.add_heading("2. Data and assumptions", level=1)
    d = ctx["dataset"]
    _para(doc, f"{d['observations']:,} {d['frequency']} returns for {len(assets)} assets from {ctx['window'][0]} to {ctx['window'][1]}. "
               f"Expected returns and risk: {ctx['assumption_label']}.")
    lower, upper = ctx["constraints"]["lower"], ctx["constraints"]["upper"]
    _table(doc, ["Asset", "Expected return", "Volatility", "Min weight", "Max weight"],
           [[assets[i], pct(ctx["mu"][i]), pct(ctx["vol"][i]), pct(lower[i], 0), pct(upper[i], 0)] for i in range(len(assets))])
    if ctx["constraints"].get("groups"):
        _bullets(doc, [f"Group limit {g.get('name') or ''}: {', '.join(assets[i] for i in g['assets'])} "
                       f"between {pct(g.get('min') or 0, 0)} and {pct(1 if g.get('max') is None else g['max'], 0)}."
                       for g in ctx["constraints"]["groups"]])
    for note in ctx.get("notes", []):
        _para(doc, note, "Note:", size=9)

    # 3 portfolios
    doc.add_heading("3. Portfolios considered", level=1)
    doc.add_picture(chart_frontier(ctx), width=Inches(6.2))
    doc.add_picture(chart_weights(ctx), width=Inches(6.2))
    _table(doc, ["Portfolio"] + assets + ["E(R)", "Risk", "Sharpe"],
           [[p["name"]] + [pct(w, 0) for w in p["weights"]] + [pct(p["exp_return"]), pct(p["volatility"]), num(p["sharpe"])] for p in ctx["portfolios"]])

    # 4 backtest
    doc.add_heading("4. Historical performance", level=1)
    b = ctx["backtest"]
    _para(doc, f"Each portfolio applied to actual returns, {ctx['rebalance_label']}, trading cost {ctx['cost_bps']:g} basis points.")
    doc.add_picture(chart_lines(b["dates"], [(p["name"], p["wealth"]) for p in b["portfolios"]], "Growth of 1"), width=Inches(6.2))
    doc.add_picture(chart_lines(b["dates"], [(p["name"], p["drawdown"]) for p in b["portfolios"]], "Drawdown", percent=True, height=2.3), width=Inches(6.2))
    head = ["Portfolio", "Return", "Risk", "Sharpe", "Sortino", "Max drawdown", "CVaR 95%"]
    rows = [[n, pct(m["ann_return"]), pct(m["volatility"]), num(m["sharpe"]), num(m["sortino"]), pct(m["max_drawdown"]), pct(m["cvar_95"])] for n, m in bt.items()]
    if ctx.get("benchmark"):
        head += ["Active", "Tracking error"]
        rows = [r + (["benchmark", ""] if bt[r[0]].get("tracking_error") is None else [pct(bt[r[0]]["active_return"]), pct(bt[r[0]]["tracking_error"])]) for r in rows]
    _table(doc, head, rows)
    at = ctx.get("attribution")
    if at:
        doc.add_heading("Where the return came from", level=2)
        head = ["Asset", "Weight", "Asset return", "Contribution"]
        rows = [[assets[i], pct(at["weights"][i], 0), pct(at["asset_return"][i]), pct(at["contribution"][i])] for i in range(len(assets))]
        total = ["Total", "100%", "", pct(at["total_return"])]
        if "allocation_effect" in at:
            head += ["Benchmark weight", "Allocation effect"]
            rows = [r + [pct(at["benchmark_weights"][i], 0), pct(at["allocation_effect"][i])] for i, r in enumerate(rows)]
            total += ["100%", pct(at["active_return"])]
        _table(doc, head, rows + [total])

    # 5 stress
    doc.add_heading("5. Stress tests", level=1)
    if ok:
        _table(doc, ["Window", "Period", "Annualised return", "Risk", "Cumulative", "Max drawdown"],
               [[w["name"], f"{w['start']} to {w['end']}", pct(w["mean_ann_return"]), pct(w["volatility"]), pct(w["cumulative"]), pct(w["max_drawdown"])] for w in ok])
    else:
        _para(doc, "The uploaded history does not cover a recognised crisis period, and no custom window was set.")

    # 6 simulation
    doc.add_heading("6. Simulated outlook", level=1)
    method = {"normal": "normally distributed returns", "student_t": "fat-tailed (Student-t) returns", "bootstrap": "resampled blocks of the actual history"}[mc["method"]]
    _para(doc, f"{mc['trials']:,} simulated paths over {mc['years']:.0f} years using {method}.")
    doc.add_picture(chart_fan(mc["fan"]["wealth"], "Years", "Growth of 1", per_year=ctx["dataset"]["periods_per_year"]), width=Inches(6.2))
    _table(doc, ["Statistic", "Value"], [
        ["Median annual return", pct(mc["median_ann_return"])], ["5th percentile annual return", pct(mc["p05_ann_return"])],
        ["Average of the worst 5% (CVaR)", pct(mc["cvar_95_ann_return"])], ["Probability of a loss", pct(mc["prob_loss"])],
        ["Median worst drawdown", pct(mc["median_max_drawdown"])], ["Ending value of 1 (5th / median / 95th)", f"{mc['p05_terminal']:.2f} / {mc['median_terminal']:.2f} / {mc['p95_terminal']:.2f}"]])

    # 7 spending
    if sp:
        doc.add_heading("7. Spending policy", level=1)
        _para(doc, f"Rule: {sp['rule_label']} at {pct(sp['rate'])}. Inflation {pct(sp['inflation'])} a year, contributions {pct(sp['contribution'])} "
                   f"of the starting value a year, costs {pct(sp['fee'], 2)}. Values are in today's money, starting at 100.")
        doc.add_picture(chart_fan(sp["value_fan"]["values"], "Years", "Fund value in today's money (start = 100)"), width=Inches(6.2))
        _table(doc, ["Statistic", "Value"], [
            ["Probability of preserving real value", pct(sp["prob_preserve"], 0)],
            ["Probability of falling below 70% of today's value at any time", pct(sp["prob_below_70"], 0)],
            ["Median real value at the end", num(sp["median_terminal_real"], 1)],
            ["Median annual spending in today's money", num(sp["median_spending_real"], 2)],
            ["Probability of a real spending cut above 10% in any year", pct(sp["prob_spending_cut_10"], 0)]])
        if sp.get("comparison"):
            _table(doc, ["Spending rate", "Preserves real value", "Falls below 70%", "Median end value", "Median annual spending"],
                   [[pct(c["rate"]), pct(c["prob_preserve"], 0), pct(c["prob_below_70"], 0), num(c["median_terminal_real"], 1), num(c["median_spending_real"], 2)] for c in sp["comparison"]])

    doc.add_heading("Method and limits", level=1)
    _bullets(doc, [
        "Portfolios are optimised on expected return and covariance (mean-variance), subject to the weight and group limits shown.",
        "Backtests and stress tests apply fixed target weights to historical returns. Past returns do not predict future returns.",
        "Simulations assume the stated return and risk hold over the horizon; real markets have regime changes that these models do not capture.",
        "This report is analytical output for discussion. It is not investment advice.",
    ])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
