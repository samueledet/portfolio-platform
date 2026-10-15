"""API tests for views, CVaR, attribution, currency, spending and the committee report."""
import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from pcp.api import app

client = TestClient(app)
SAA = [0.1, 0.1, 0.5, 0.1, 0.1, 0.1]
B6040 = [0.6, 0, 0, 0, 0.4, 0]


@pytest.fixture(scope="module")
def ds():
    return client.post("/api/datasets/sample/indices").json()


def post(ds, path, **body):
    res = client.post(f"/api/datasets/{ds['id']}/{path}", json=body)
    assert res.status_code == 200, res.text
    return res.json()


def test_custom_correlation_blend_and_views(ds):
    base = post(ds, "statistics")
    corr = [[1.0 if i == j else 0.3 for j in range(6)] for i in range(6)]
    custom = post(ds, "statistics", assumptions={"source": "custom", "corr": corr})
    assert custom["assumptions"]["corr"][0][1] == pytest.approx(0.3) and custom["assumptions"]["notes"] == []
    bad = [[1.0 if i == j else -0.9 for j in range(6)] for i in range(6)]
    repaired = post(ds, "statistics", assumptions={"source": "custom", "corr": bad})
    assert "nearest valid matrix" in repaired["assumptions"]["notes"][0]
    mu, vol = base["historical"]["mu"], base["historical"]["vol"]
    blend = post(ds, "statistics", assumptions={"source": "blend", "sources": [
        {"name": "House", "weight": 1, "mu": mu, "vol": vol}, {"name": "Manager", "weight": 3, "mu": [m + 0.04 for m in mu], "vol": vol}]})
    assert blend["assumptions"]["mu"][0] == pytest.approx(mu[0] + 0.03)
    bl = post(ds, "statistics", rf=0.02, assumptions={"source": "black_litterman", "market_weights": [1] * 6,
                                                      "views": [{"type": "absolute", "asset": 5, "value": 0.15, "confidence": 0.9}]})
    assert bl["assumptions"]["mu"][5] > bl["assumptions"]["prior_mu"][5]
    err = client.post(f"/api/datasets/{ds['id']}/statistics", json={"assumptions": {"source": "blend", "sources": []}})
    assert err.status_code == 422 and "forecast set" in err.json()["detail"]


def test_min_cvar_attribution_and_spending(ds):
    p = post(ds, "optimise", objective="min_cvar", beta=0.95, constraints={"upper": [0.5] * 6})
    assert sum(p["weights"]) == pytest.approx(1) and max(p["weights"]) <= 0.5 + 1e-6 and p["period_cvar"] > 0
    at = post(ds, "attribution", weights=SAA, benchmark_weights=B6040, start="2008-01-01", end="2017-12-31")
    assert sum(at["allocation_effect"]) == pytest.approx(at["active_return"], abs=1e-9)
    assert sum(at["risk_contribution"]) == pytest.approx(1, abs=1e-6)
    sp = post(ds, "spending", weights=SAA, years=20, trials=1000, rule="hybrid", rate=0.045, compare_rates=[0.03, 0.05])
    assert 0 <= sp["prob_preserve"] <= 1 and len(sp["value_fan"]["values"]) == 5 and len(sp["value_fan"]["values"][0]) == 21
    assert sp["comparison"][0]["prob_preserve"] >= sp["comparison"][1]["prob_preserve"]
    bad = client.post(f"/api/datasets/{ds['id']}/spending", json={"weights": SAA, "rule": "nope"})
    assert bad.status_code == 422


def test_currency_endpoint_on_fx_sample():
    fx = client.post("/api/datasets/sample/fx").json()
    h = client.post(f"/api/datasets/{fx['id']}/currency", json={"asset": 0, "currency": 1, "hedge_ratio": 0.5, "base_rate": 0.02684, "foreign_rate": 0.01219}).json()
    assert h["metrics"]["Unhedged"]["cumulative"] == pytest.approx(-0.13893, abs=1e-4) and len(h["curve"]) == 13
    same = client.post(f"/api/datasets/{fx['id']}/currency", json={"asset": 0, "currency": 0})
    assert same.status_code == 422


def test_committee_report_is_a_complete_word_document(ds):
    ports = [{"name": "Proposed SAA", "weights": SAA}, {"name": "60/40", "weights": B6040}, {"name": "Equal weight", "weights": [1] * 6}]
    res = client.post(f"/api/datasets/{ds['id']}/report.docx", json={
        "portfolios": ports, "focus": "Proposed SAA", "benchmark": "60/40", "client": "Example Endowment", "rf": 0.02,
        "constraints": {"lower": [0.05] * 6, "groups": [{"name": "Equities", "assets": [0, 1, 2], "max": 0.7}]},
        "simulation": {"years": 10, "trials": 800}, "spending": {"rule": "smoothed", "rate": 0.04, "years": 20, "trials": 800, "compare_rates": [0.03, 0.04, 0.05]}})
    assert res.status_code == 200, res.text
    doc = Document(io.BytesIO(res.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    for heading in ("Summary for the committee", "Data and assumptions", "Portfolios considered", "Historical performance",
                    "Stress tests", "Simulated outlook", "Spending policy", "Method and limits"):
        assert heading in text
    assert "Proposed SAA" in text and "Example Endowment" in text and "not investment advice" in text
    assert len(doc.inline_shapes) == 6 and len(doc.tables) >= 7          # charts and tables made it in
    plain = client.post(f"/api/datasets/{ds['id']}/report.docx", json={"portfolios": ports[:1]})
    assert plain.status_code == 200 and "Spending policy" not in "\n".join(p.text for p in Document(io.BytesIO(plain.content)).paragraphs)


# ------------------------------------------------------------------ mandates through the API
@pytest.mark.parametrize("mandate", [
    {"kind": "long_short", "net": 1.0, "gross": 1.6, "short_cost": 0.005},
    {"kind": "market_neutral", "gross": 2.0},
    {"kind": "leveraged", "leverage": 1.5, "borrow_spread": 0.01},
    {"kind": "benchmark_relative", "benchmark": [0.3, 0.1, 0.1, 0.1, 0.3, 0.1], "active_limit": 0.1},
    {"kind": "surplus", "liability": 4, "funding_ratio": 1.1},
])
def test_every_step_runs_under_each_mandate(ds, mandate):
    common = {"mandate": mandate, "rf": 0.02}
    f = post(ds, "frontier", **common)
    assert len(f["points"]) > 5 and f["axes"]["x"] and f["mandate"]["kind"] == mandate["kind"]
    ports = [{"name": k, "weights": v["weights"]} for k, v in f["portfolios"].items()]
    bt = post(ds, "backtest", portfolios=ports, benchmark=ports[-1]["name"], **common)
    assert all(p["metrics"]["periods"] == 215 for p in bt["portfolios"])
    w = ports[0]["weights"]
    assert post(ds, "stress", weights=w, **common)["windows"]
    assert 0 <= post(ds, "simulate", weights=w, years=2, trials=300, **common)["prob_loss"] <= 1
    assert post(ds, "spending", weights=w, years=10, trials=300, **common)["years"] == 10
    at = post(ds, "attribution", weights=w, benchmark_weights=ports[-1]["weights"], **common)
    assert sum(at["contribution"]) == pytest.approx(at["total_return"], abs=1e-9)
    res = client.post(f"/api/datasets/{ds['id']}/report.docx", json={"portfolios": ports, "focus": ports[0]["name"], **common,
                                                                     "simulation": {"years": 2, "trials": 300}})
    assert res.status_code == 200, res.text
    assert "Mandate:" in "\n".join(p.text for p in Document(io.BytesIO(res.content)).paragraphs)


def test_market_neutral_backtest_keeps_its_cash_and_shorts(ds):
    w = [0.3, -0.3, 0.0, 0.0, 0.2, -0.2]                     # net zero: 100% cash collateral
    bt = post(ds, "backtest", portfolios=[{"name": "MN", "weights": w}], mandate={"kind": "market_neutral"}, rf=0.03)
    plain = post(ds, "backtest", portfolios=[{"name": "MN", "weights": w}], rf=0.03, mandate={"kind": "market_neutral", "short_cost": 0.0})
    assert bt["portfolios"][0]["metrics"]["ann_return"] == pytest.approx(plain["portfolios"][0]["metrics"]["ann_return"])
    bad = client.post(f"/api/datasets/{ds['id']}/optimise", json={"objective": "min_variance", "constraints": {"lower": [-0.2] * 6}})
    assert bad.status_code == 422 and "short positions" in bad.json()["detail"]
